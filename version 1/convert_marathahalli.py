import csv
import openpyxl
from openpyxl.styles import Font

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Marathahalli IT Jobs"

with open('Marathahalli_IT_Jobs.csv', 'r', encoding='utf-8') as f:
    reader = csv.reader(f)
    for row_idx, row in enumerate(reader, start=1):
        ws.append(row)
        # Format the 5th column as a clickable hyperlink if it's not the header
        if row_idx > 1 and len(row) >= 5:
            cell = ws.cell(row=row_idx, column=5)
            cell.hyperlink = row[4]
            cell.font = Font(color="0563C1", underline="single")

# Make header bold
for cell in ws[1]:
    cell.font = Font(bold=True)

# Auto-adjust column widths
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if cell.value and len(str(cell.value)) > max_length:
                max_length = len(cell.value)
        except:
            pass
    ws.column_dimensions[column].width = min(max_length + 2, 50)

wb.save('Marathahalli_IT_Jobs.xlsx')
